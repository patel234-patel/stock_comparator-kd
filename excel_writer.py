"""
excel_writer.py — Writes live price data to Excel
===================================================
Creates a new .xlsx file each day. Auto-saves every N rows.
"""

import os
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import OUTPUT_DIR, SAVE_EVERY


HEADERS = [
    "Timestamp",
    "Angel One Price (₹)",
    "Motilal Price (₹)",
    "Difference (₹)",
    "Diff (%)",
]

# Column widths
COL_WIDTHS = [22, 20, 20, 18, 12]

# Colors
HDR_BG    = "1F4E79"   # dark blue
HDR_FG    = "FFFFFF"   # white
POS_FILL  = "C6EFCE"   # green  — Angel > Motilal
NEG_FILL  = "FFC7CE"   # red    — Angel < Motilal
ZERO_FILL = "FFEB9C"   # yellow — equal


class ExcelWriter:
    def __init__(self, symbol: str):
        self.symbol   = symbol
        self._rows    = []
        self._counter = 0

        os.makedirs(OUTPUT_DIR, exist_ok=True)
        date_str = datetime.date.today().strftime("%Y-%m-%d")
        filename = f"{symbol}_prices_{date_str}.xlsx"
        self.filepath = os.path.join(OUTPUT_DIR, filename)

        self._wb = openpyxl.Workbook()
        self._ws = self._wb.active
        self._ws.title = f"{symbol} Prices"

        self._write_headers()

    # ── Internal helpers ───────────────────────────────────────────────────────
    def _write_headers(self):
        ws = self._ws

        # Title row
        ws.merge_cells("A1:E1")
        title_cell = ws["A1"]
        title_cell.value = f"{self.symbol} — Angel One vs Motilal Oswal Live Prices"
        title_cell.font      = Font(name="Arial", bold=True, size=13, color="FFFFFF")
        title_cell.fill      = PatternFill("solid", fgColor="0D3349")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Header row (row 2)
        thin = Side(style="thin", color="AAAAAA")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
            cell = ws.cell(row=2, column=col, value=header)
            cell.font      = Font(name="Arial", bold=True, color=HDR_FG, size=10)
            cell.fill      = PatternFill("solid", fgColor=HDR_BG)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = border
            ws.column_dimensions[get_column_letter(col)].width = width

        ws.row_dimensions[2].height = 20
        ws.freeze_panes = "A3"   # freeze title + header

    def _row_fill(self, diff: float) -> str | None:
        if diff > 0:   return POS_FILL
        if diff < 0:   return NEG_FILL
        return ZERO_FILL

    # ── Public API ─────────────────────────────────────────────────────────────
    def append(self, row: dict):
        """Add one second's data. Auto-saves every SAVE_EVERY rows."""
        self._rows.append(row)
        self._counter += 1

        ws   = self._ws
        r    = len(self._rows) + 2          # +2 for title + header rows
        fill = PatternFill("solid", fgColor=self._row_fill(row["difference"]))

        thin   = Side(style="thin", color="DDDDDD")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        values = [
            row["timestamp"].strftime("%Y-%m-%d %H:%M:%S"),
            row["angel_price"],
            row["motilal_price"],
            row["difference"],
            row["diff_pct"],
        ]
        number_formats = [None, "#,##0.00", "#,##0.00", "#,##0.0000", "0.0000%"]

        for col, (val, fmt) in enumerate(zip(values, number_formats), start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font      = Font(name="Arial", size=9)
            cell.fill      = fill
            cell.border    = border
            cell.alignment = Alignment(horizontal="center")
            if fmt and isinstance(val, (int, float)):
                cell.number_format = fmt

        if self._counter % SAVE_EVERY == 0:
            self.save()

    def save(self):
        self._wb.save(self.filepath)
